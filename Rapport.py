import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

DB_NAME = "fruteria.db"

def obtener_constantes_frutas():
    from Fruits import BOLSA_INFO, FRUTAS, TIPO_IVA
    return FRUTAS, BOLSA_INFO, TIPO_IVA

def inicializar_db():
    FRUTAS, BOLSA_INFO, _ = obtener_constantes_frutas()
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS stock (
            clave TEXT PRIMARY KEY,
            nombre TEXT,
            unidad TEXT,
            cantidad REAL
        )
    ''')
    
    cursor.execute("SELECT COUNT(*) FROM stock")
    if cursor.fetchone()[0] == 0:
        for clave, data in FRUTAS.items():
            nombre, unidad = data[1], data[3]
            stock_inicial = 50.0 if unidad == "kg" else 30.0
            cursor.execute("INSERT INTO stock VALUES (?, ?, ?, ?)", (clave, nombre, unidad, stock_inicial))
        cursor.execute("INSERT INTO stock VALUES (?, ?, ?, ?)", ("bolsa", BOLSA_INFO[1], BOLSA_INFO[3], 500.0))
    
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS ventas (
            id_ticket TEXT PRIMARY KEY,
            fecha_hora TEXT,
            metodo_pago TEXT,
            total REAL,
            base REAL,
            iva REAL,
            entrega REAL,
            cambio REAL
        )
    ''')
    
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS detalle_ventas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            id_ticket TEXT,
            clave TEXT,
            nombre TEXT,
            cantidad REAL,
            unidad TEXT,
            pvp REAL,
            subtotal REAL,
            FOREIGN KEY (id_ticket) REFERENCES ventas(id_ticket)
        )
    ''')
    
    conn.commit()
    conn.close()

def registrar_venta_db(id_ticket, cesta, entrega, cambio, metodo_pago):
    _, _, TIPO_IVA = obtener_constantes_frutas()
    inicializar_db()
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    total = sum(item["total"] for item in cesta.values())
    base = total / (1 + TIPO_IVA)
    iva = total - base
    fecha_hora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    cursor.execute('''
        INSERT INTO ventas VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    ''', (id_ticket, fecha_hora, metodo_pago, total, base, iva, entrega, cambio))
    
    for clave, item in cesta.items():
        cursor.execute('''
            INSERT INTO detalle_ventas (id_ticket, clave, nombre, cantidad, unidad, pvp, subtotal)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        ''', (id_ticket, clave, item['nombre'], item['kg'], item['unidad'], item['pvp'], item['total']))
        
        cursor.execute("UPDATE stock SET cantidad = cantidad - ? WHERE clave = ?", (item['kg'], clave))
        
    conn.commit()
    conn.close()

def aplicar_estilos_generales(ws, titulo):
    fill_header = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    
    border_thin = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )
    
    ws.merge_cells("A1:F1")
    ws["A1"] = f"FRUTERÍA L'HORT - {titulo.upper()}"
    ws["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="2E7D32")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    
    ws["A2"] = f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    ws["A2"].font = Font(name="Segoe UI", size=9, italic=True, color="666666")
    
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = border_thin
            if cell.row == 4:
                cell.fill = fill_header
                cell.font = font_header
                cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

def generar_reporte_ventas(ruta_destino=None):
    inicializar_db()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte de Ventas"
    
    headers = ["Nº Ticket", "Fecha / Hora", "Método Pago", "Base (€)", "IVA (€)", "Total (€)"]
    ws.append([])
    ws.append(headers)
    
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute("SELECT id_ticket, fecha_hora, metodo_pago, base, iva, total FROM ventas ORDER BY fecha_hora DESC")
    filas = cursor.fetchall()
    conn.close()
    
    for r in filas:
        ws.append([r[0], r[1], r[2], round(r[3], 2), round(r[4], 2), round(r[5], 2)])
        
    ult_fila = len(filas) + 4
    if filas:
        ws.append(["TOTALES", "", "", f"=SUM(D5:D{ult_fila})", f"=SUM(E5:E{ult_fila})", f"=SUM(F5:F{ult_fila})"])
        for col in range(1, 7):
            ws.cell(row=ws.max_row, column=col).font = Font(bold=True)

    aplicar_estilos_generales(ws, "Reporte de Ventas")
    filename = ruta_destino if ruta_destino else f"reporte_ventas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(filename)
    return filename

def generar_reporte_stock(ruta_destino=None):
    inicializar_db()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte de Stock"
    
    headers = ["Código", "Producto", "Unidad", "Stock Disponible", "Estado"]
    ws.append([])
    ws.append(headers)
    
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute("SELECT clave, nombre, unidad, cantidad FROM stock ORDER BY nombre ASC")
    filas = cursor.fetchall()
    conn.close()
    
    fill_alerta = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")
    
    for r in filas:
        estado = "NORMAL" if r[3] > 10.0 else "STOCK BAJO"
        ws.append([r[0], r[1], r[2], round(r[3], 2), estado])
        
        if estado == "STOCK BAJO":
            for col in range(1, 6):
                ws.cell(row=ws.max_row, column=col).fill = fill_alerta

    aplicar_estilos_generales(ws, "Control de Stock")
    filename = ruta_destino if ruta_destino else f"reporte_stock_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(filename)
    return filename

def generar_reporte_cuadre_caja(ruta_destino=None):
    inicializar_db()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cuadre de Caja"
    
    headers = ["Método de Pago", "Nº Operaciones", "Total Recaudado (€)"]
    ws.append([])
    ws.append(headers)
    
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute("""
        SELECT metodo_pago, COUNT(*), SUM(total)
        FROM ventas
        WHERE DATE(fecha_hora) = DATE('now', 'localtime')
        GROUP BY metodo_pago
    """)
    filas = cursor.fetchall()
    conn.close()
    
    for r in filas:
        ws.append([r[0], r[1], round(r[2], 2)])
        
    ult_fila = len(filas) + 4
    if filas:
        ws.append(["TOTAL DÍA", f"=SUM(B5:B{ult_fila})", f"=SUM(C5:C{ult_fila})"])
        for col in range(1, 4):
            ws.cell(row=ws.max_row, column=col).font = Font(bold=True)

    aplicar_estilos_generales(ws, "Cuadre de Caja Diario")
    filename = ruta_destino if ruta_destino else f"reporte_cuadre_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(filename)
    return filename